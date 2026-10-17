"""
Gera public/materiais/modelo-de-impugnacao-de-edital.docx.

Por que um gerador, e não o .docx solto no repositório: o arquivo binário não
mostra diff. Quem revisa uma mudança no modelo precisa ver o texto que mudou,
e é este arquivo que mostra. Rodar de novo depois de editar:

    pip install python-docx
    python3 materiais/gerar_modelo_de_impugnacao.py

O que o modelo pode e não pode dizer
------------------------------------
Ele organiza o pedido; não argumenta por ninguém. Tudo que é fato de lei está
escrito e conferido no texto da Lei 14.133 no Planalto em 17/10/2026:

- art. 164, caput: qualquer pessoa pode impugnar, protocolo até 3 dias úteis
  antes da abertura do certame;
- art. 164, parágrafo único: resposta em até 3 dias úteis, limitada ao último
  dia útil anterior à abertura, divulgada em sítio eletrônico oficial;
- art. 183: contagem com exclusão do dia do começo e inclusão do vencimento;
- art. 55, § 1º: modificação do edital exige nova divulgação e reabertura dos
  prazos, salvo quando não comprometer a formulação das propostas.

O fundamento de cada caso (o dispositivo violado, a súmula, o acórdão) fica em
campo para preencher. Escrever argumento pronto seria dar parecer, e o site não
dá parecer.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

SAIDA = Path(__file__).resolve().parent.parent / "public" / "materiais" / "modelo-de-impugnacao-de-edital.docx"

CINZA = RGBColor(0x55, 0x5F, 0x6D)


def paragrafo(doc, texto, *, negrito=False, italico=False, cor=None, tamanho=None, alinhar=None, depois=6):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = negrito
    run.italic = italico
    if cor:
        run.font.color.rgb = cor
    if tamanho:
        run.font.size = Pt(tamanho)
    if alinhar:
        p.alignment = alinhar
    p.paragraph_format.space_after = Pt(depois)
    return p


def campo(doc, texto):
    """Instrução de preenchimento: entre colchetes e em cinza, para saltar aos olhos."""
    return paragrafo(doc, texto, italico=True, cor=CINZA)


def gerar():
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    paragrafo(doc, "MODELO DE IMPUGNAÇÃO DE EDITAL", negrito=True, tamanho=14, alinhar=WD_ALIGN_PARAGRAPH.CENTER, depois=2)
    paragrafo(doc, "Lei nº 14.133/2021, art. 164", alinhar=WD_ALIGN_PARAGRAPH.CENTER, cor=CINZA, depois=12)

    paragrafo(
        doc,
        "Antes de usar (apague este quadro antes de enviar): este modelo organiza o pedido, "
        "mas não é parecer jurídico e não substitui a análise de um advogado. O edital "
        "prevalece. Confira nele o meio de protocolo (sistema, e-mail ou outro), o "
        "horário-limite e se ele pede documento de representação junto com a impugnação. "
        "Guarde o comprovante de envio com data e hora e acompanhe a resposta e uma "
        "eventual republicação. Apague também todos os textos entre colchetes.",
        italico=True,
        cor=CINZA,
        tamanho=10,
        depois=14,
    )

    paragrafo(doc, "Ao(À) [Agente de Contratação / Pregoeiro(a) / Comissão de Contratação]")
    paragrafo(doc, "[Nome do órgão ou da entidade]", depois=12)

    paragrafo(
        doc,
        "Referência: [Pregão Eletrônico / Concorrência / outra modalidade] nº [   ] · "
        "Processo nº [   ] · Abertura da sessão: [data e hora]",
        negrito=True,
        depois=12,
    )

    paragrafo(
        doc,
        "[RAZÃO SOCIAL], inscrita no CNPJ sob o nº [   ], com sede em [endereço completo], "
        "neste ato representada por [nome do representante legal], [cargo], vem, com "
        "fundamento no art. 164 da Lei nº 14.133/2021, apresentar IMPUGNAÇÃO ao edital em "
        "referência, pelos fatos e fundamentos a seguir.",
        depois=12,
    )

    paragrafo(doc, "1. Da tempestividade", negrito=True)
    paragrafo(
        doc,
        "O art. 164 da Lei nº 14.133/2021 admite a impugnação protocolada até 3 (três) dias "
        "úteis antes da data de abertura do certame. A abertura está marcada para [data]. "
        "Contado na forma do art. 183 da mesma Lei, com exclusão do dia do começo e inclusão "
        "do dia do vencimento, o prazo termina em [data-limite]. Esta impugnação é "
        "protocolada em [data do protocolo], dentro do prazo.",
    )
    campo(
        doc,
        "[Para contar: exclua o dia da abertura e volte 3 dias úteis do órgão que licita. "
        "Exemplo: sessão numa terça, sem feriado no caminho, último dia na quinta anterior. "
        "Se o edital ou o sistema informar a data-limite, use a dele.]",
    )

    paragrafo(doc, "2. Da cláusula impugnada", negrito=True)
    paragrafo(
        doc,
        "O item [número] do [edital / termo de referência / anexo], na página [   ], "
        "estabelece:",
    )
    campo(doc, "\"[Transcreva o texto exato da cláusula, entre aspas.]\"")

    paragrafo(doc, "3. Dos fundamentos", negrito=True)
    campo(
        doc,
        "[Explique de forma objetiva por que a exigência contraria a Lei nº 14.133/2021 ou "
        "restringe a competição sem justificativa. Indique o dispositivo violado: artigo, "
        "parágrafo e inciso.]",
    )
    campo(
        doc,
        "[Se houver, cite a súmula ou o acórdão do Tribunal de Contas da União sobre o mesmo "
        "ponto, com número e ano, e transcreva o trecho que se aplica ao seu caso.]",
    )
    campo(
        doc,
        "[Se a exigência tem efeito concreto, descreva-o com números: quantas empresas ficam "
        "de fora, que custo desproporcional ela cria, por que o prazo é inexequível.]",
    )

    paragrafo(doc, "4. Do pedido", negrito=True)
    paragrafo(doc, "Diante do exposto, requer:")
    paragrafo(doc, "a) o recebimento e o processamento desta impugnação;")
    paragrafo(
        doc,
        "b) a [alteração / supressão] do item [número], para que passe a constar: "
        "\"[redação proposta]\";",
    )
    paragrafo(
        doc,
        "c) a nova divulgação do edital alterado, na mesma forma da divulgação inicial, com a "
        "reabertura dos prazos, caso a alteração comprometa a formulação das propostas "
        "(art. 55, § 1º, da Lei nº 14.133/2021);",
    )
    paragrafo(
        doc,
        "d) a divulgação da resposta em sítio eletrônico oficial no prazo de até 3 (três) "
        "dias úteis, limitado ao último dia útil anterior à data da abertura do certame "
        "(art. 164, parágrafo único, da Lei nº 14.133/2021).",
        depois=14,
    )

    paragrafo(doc, "Nestes termos, pede deferimento.", depois=14)
    paragrafo(doc, "[Cidade], [data].", depois=28)
    paragrafo(doc, "______________________________________", depois=2)
    paragrafo(doc, "[Nome do representante legal]", depois=2)
    paragrafo(doc, "[Cargo] · [Razão social] · CNPJ [   ]", depois=0)

    # Nada depois da assinatura: o que ficasse aqui iria parar no protocolo.
    # A origem do modelo vai nas propriedades do arquivo, não no corpo.
    doc.core_properties.comments = (
        "Modelo gratuito do Licitante Vencedor: "
        "licitantevencedor.com.br/blog/prazo-para-impugnar-edital-de-licitacao/"
    )
    doc.core_properties.title = "Modelo de impugnação de edital (Lei 14.133/2021, art. 164)"
    doc.core_properties.author = "Licitante Vencedor"

    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SAIDA)
    print(f"gerado: {SAIDA}")


if __name__ == "__main__":
    gerar()
